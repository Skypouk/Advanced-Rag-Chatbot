"""Document extraction with character offsets and (for PDFs) word geometry.

Every extractor returns an :class:`ExtractedDocument`: a list of pages, each
holding the exact text the viewer displays. For PDFs, each page also carries
*spans*, the word boxes and their character offsets in that text, which lets a
sentence's ``[start, end)`` offsets be mapped to highlight rectangles on the
rendered page.

Strategy per file type:

* **PDF**: PyMuPDF word extraction (exact per-word boxes). If the PDF has no
  text layer (scanned), Docling with OCR is used and each highlight becomes
  the box of the paragraph that contains the sentence.
* **DOCX**: Docling, falling back to python-docx. The text is split into
  virtual pages.
* **TXT**: read directly and split into virtual pages.
"""

import logging
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import TYPE_CHECKING, Dict, List, Optional, Tuple

from config import Settings, get_settings

if TYPE_CHECKING:  # pragma: no cover - imported for type hints only
    from docling.document_converter import DocumentConverter

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}

# Below this many characters per page on average, a PDF is treated as scanned.
_MIN_CHARS_PER_PAGE = 20

Rect = Tuple[float, float, float, float]


class ExtractionError(Exception):
    """Raised when no text could be extracted from a document."""


class UnsupportedFileTypeError(ExtractionError):
    """Raised when a file extension is not supported."""


@dataclass
class TextSpan:
    """A run of page text with its bounding box (PDF points, top-left origin)."""

    start: int
    end: int
    rect: Rect
    # Spans sharing a line key are merged into one highlight rectangle.
    line: str


@dataclass
class ExtractedPage:
    """One page of extracted text."""

    number: int  # 1-based
    text: str
    width: Optional[float] = None
    height: Optional[float] = None
    spans: List[TextSpan] = field(default_factory=list)


@dataclass
class ExtractedDocument:
    """Result of extracting a document."""

    file_type: str  # "pdf" | "docx" | "txt"
    extractor: str
    pages: List[ExtractedPage]

    @property
    def text(self) -> str:
        """Return the full document text (pages separated by blank lines)."""
        return "\n\n".join(page.text for page in self.pages)


# ── PDF: PyMuPDF ────────────────────────────────────────────────────────────


def _extract_pdf_with_pymupdf(path: Path) -> ExtractedDocument:
    """Extract page text plus per-word boxes from a PDF using PyMuPDF.

    Words on the same line are joined with a space, lines of the same block
    with a space (so a sentence wrapped over lines stays one sentence) and
    blocks with a blank line (a paragraph boundary for the segmenter).
    """
    import pymupdf

    pages: List[ExtractedPage] = []
    with pymupdf.open(str(path)) as pdf:
        if pdf.needs_pass:
            raise ExtractionError("The PDF is password protected.")
        for number, page in enumerate(pdf, start=1):
            # Word boxes are reported in unrotated page space; map them to
            # the displayed (rotated) page that the viewer renders.
            matrix = page.rotation_matrix if page.rotation else None
            parts: List[str] = []
            spans: List[TextSpan] = []
            position = 0
            previous: Optional[Tuple[int, int]] = None
            for x0, y0, x1, y1, word, block, line, _ in page.get_text("words"):
                if previous is not None:
                    separator = " " if block == previous[0] else "\n\n"
                    parts.append(separator)
                    position += len(separator)
                rect = pymupdf.Rect(x0, y0, x1, y1)
                if matrix is not None:
                    rect = rect * matrix
                spans.append(
                    TextSpan(
                        start=position,
                        end=position + len(word),
                        rect=(rect.x0, rect.y0, rect.x1, rect.y1),
                        line=f"{block}:{line}",
                    )
                )
                parts.append(word)
                position += len(word)
                previous = (block, line)
            pages.append(
                ExtractedPage(
                    number=number,
                    text="".join(parts),
                    width=page.rect.width,
                    height=page.rect.height,
                    spans=spans,
                )
            )
    return ExtractedDocument(file_type="pdf", extractor="pymupdf", pages=pages)


# ── Docling (DOCX + OCR for scanned PDFs) ──────────────────────────────────


@lru_cache(maxsize=1)
def _get_docling_converter() -> "DocumentConverter":
    """Create the Docling converter once; model loading is expensive."""
    from docling.document_converter import DocumentConverter

    return DocumentConverter()


def _docling_item_text(item: object, doc: object) -> str:
    """Return the text of a Docling item (tables are exported as Markdown)."""
    text = getattr(item, "text", None)
    if isinstance(text, str):
        return text
    export = getattr(item, "export_to_markdown", None)
    if callable(export):
        try:
            return str(export(doc=doc))
        except Exception:  # noqa: BLE001 - skip items that cannot be exported
            return ""
    return ""


def _extract_with_docling(path: Path, file_type: str) -> ExtractedDocument:
    """Convert a document with Docling, keeping page numbers and item boxes.

    Each Docling item (paragraph, heading, list item, table) becomes one span,
    so its provenance box is used as the highlight area.
    """
    result = _get_docling_converter().convert(str(path))
    doc = result.document

    page_sizes: Dict[int, Tuple[float, float]] = {
        int(no): (float(page.size.width), float(page.size.height)) for no, page in doc.pages.items()
    }
    parts: Dict[int, List[str]] = {no: [] for no in page_sizes}
    lengths: Dict[int, int] = {no: 0 for no in page_sizes}
    spans: Dict[int, List[TextSpan]] = {no: [] for no in page_sizes}

    for index, (item, _level) in enumerate(doc.iterate_items()):
        text = _docling_item_text(item, doc).strip()
        if not text:
            continue
        provenance = getattr(item, "prov", None) or []
        page_no = int(provenance[0].page_no) if provenance else 1
        parts.setdefault(page_no, [])
        lengths.setdefault(page_no, 0)
        spans.setdefault(page_no, [])

        if parts[page_no]:
            parts[page_no].append("\n\n")
            lengths[page_no] += 2
        start = lengths[page_no]
        parts[page_no].append(text)
        lengths[page_no] += len(text)

        if provenance and page_no in page_sizes:
            box = provenance[0].bbox.to_top_left_origin(page_height=page_sizes[page_no][1])
            spans[page_no].append(
                TextSpan(start=start, end=start + len(text), rect=(box.l, box.t, box.r, box.b), line=f"item{index}")
            )

    page_numbers = sorted(parts) or [1]
    pages = [
        ExtractedPage(
            number=no,
            text="".join(parts.get(no, [])),
            width=page_sizes.get(no, (None, None))[0],
            height=page_sizes.get(no, (None, None))[1],
            spans=spans.get(no, []),
        )
        for no in range(1, max(page_numbers) + 1)
    ]
    return ExtractedDocument(file_type=file_type, extractor="docling", pages=pages)


# ── DOCX fallback / TXT ────────────────────────────────────────────────────


def _extract_docx_with_python_docx(path: Path) -> str:
    """Extract paragraph and table text from a DOCX file."""
    import docx

    document = docx.Document(str(path))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


def _read_text_file(path: Path) -> str:
    """Read a plain-text file, tolerating common encodings."""
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="ignore")


def paginate_text(text: str, max_chars: int) -> List[ExtractedPage]:
    """Split plain text into virtual pages of at most ``max_chars``.

    Page breaks are placed at paragraph boundaries when possible, then at line
    boundaries, and only as a last resort inside a line.
    """
    text = re.sub(r"\r\n?", "\n", text).strip()
    pages: List[ExtractedPage] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            window = text[start:end]
            cut = window.rfind("\n\n")
            if cut < max_chars // 3:
                cut = window.rfind("\n")
            if cut < max_chars // 3:
                cut = window.rfind(" ")
            if cut > 0:
                end = start + cut
        page_text = text[start:end].strip()
        if page_text:
            pages.append(ExtractedPage(number=len(pages) + 1, text=page_text))
        start = end
        while start < len(text) and text[start].isspace():
            start += 1
    return pages


# ── Public API ─────────────────────────────────────────────────────────────


def _file_type(filename: str) -> str:
    """Return the normalized file type for ``filename`` or raise if unsupported."""
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{extension or 'unknown'}'. Allowed: PDF, DOCX, TXT."
        )
    return extension.lstrip(".")


def extract_document(
    file_path: str | Path, filename: Optional[str] = None, settings: Optional[Settings] = None
) -> ExtractedDocument:
    """Extract pages (with offsets and, for PDFs, word boxes) from a document.

    Args:
        file_path: Path to the file on disk.
        filename: Original filename, used for type detection and messages.
        settings: Optional settings override.

    Returns:
        The extracted document.

    Raises:
        UnsupportedFileTypeError: If the file type is not supported.
        ExtractionError: If no text could be extracted.
    """
    settings = settings or get_settings()
    path = Path(file_path)
    display_name = filename or path.name
    file_type = _file_type(display_name)

    try:
        if file_type == "pdf":
            document = _extract_pdf_with_pymupdf(path)
            total_chars = sum(len(page.text) for page in document.pages)
            if total_chars < _MIN_CHARS_PER_PAGE * max(1, len(document.pages)):
                logger.info("'%s' has no usable text layer; running Docling OCR", display_name)
                try:
                    document = _extract_with_docling(path, "pdf")
                except Exception as exc:  # noqa: BLE001
                    logger.warning("Docling OCR failed for '%s': %s", display_name, exc)
        elif file_type == "docx":
            try:
                docling_doc = _extract_with_docling(path, "docx")
                document = ExtractedDocument(
                    file_type="docx",
                    extractor="docling",
                    pages=paginate_text(docling_doc.text, settings.TEXT_PAGE_CHARS),
                )
            except Exception as exc:  # noqa: BLE001 - any Docling failure triggers fallback
                logger.warning("Docling failed for '%s' (%s); using python-docx", display_name, exc)
                document = ExtractedDocument(
                    file_type="docx",
                    extractor="python-docx",
                    pages=paginate_text(_extract_docx_with_python_docx(path), settings.TEXT_PAGE_CHARS),
                )
        else:
            document = ExtractedDocument(
                file_type="txt",
                extractor="text",
                pages=paginate_text(_read_text_file(path), settings.TEXT_PAGE_CHARS),
            )
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not read '{display_name}': {exc}") from exc

    if not any(page.text.strip() for page in document.pages):
        raise ExtractionError(
            f"No extractable text found in '{display_name}'. "
            "The file may be empty, image-only, or password protected."
        )
    logger.info(
        "Extracted '%s' with %s: %d pages, %d chars",
        display_name,
        document.extractor,
        len(document.pages),
        sum(len(page.text) for page in document.pages),
    )
    return document


def extract_text(file_path: str | Path, filename: Optional[str] = None) -> str:
    """Extract a document and return its full text as a single string."""
    return extract_document(file_path, filename).text
